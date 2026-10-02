#!/usr/bin/env python3
"""属人アカウント「いけ」（@cw_uragawa）の導線の数字を、日ごとのKPIの点にする。

読むもの:
  --dash PATH       kpi-dashboard（公開リポジトリ kazuki0806/kpi-dashboard）の clone
                      kpi_data_A.json   Threadsの投稿ごとの表示回数・フォロワー
                      line/line_A.json  公式LINEの友だち・ブロック・届く人数
                      note/note_data.json  noteの閲覧数（記事ごと・累計）
  --sheet PATH      「【いけ】属人ルート 面談・成約」を xlsx に書き出したもの（任意）
  --form-json PATH  3分チェック（Addnessフォーム）の list_form_responses の結果を
                    responses の配列のまま保存したもの（任意）

出力（--points DIR）:
  <キー>.json      record_analytics_metric_points にそのまま渡す点
  <キー>_<YYYY-MM>.parquet  証憑（月ごと。期間は月の1日〜末日で固定）
  summary.json     キーごとの点の数と、証憑の一覧（period_from・period_to・file・size_bytes・checksum・rows）

証憑の期間を月で固定しているのは、Addnessが「既にある証憑と期間がまったく同じか、
まったく重ならない証憑」しか受け付けないため。毎日入れ直しても、同じ月は同じ期間になる。
"""
import argparse
import calendar
import collections
import datetime as dt
import hashlib
import json
import os
import re
import sys

# キー → (KPI名, 種類) 種類: flow=日ごとの量 / stock=その日の残高
METRICS = {
    "threads_views": ("いけ Threads 表示回数（投稿日別）", "flow"),
    "threads_posts": ("いけ Threads 投稿数", "flow"),
    "threads_followers": ("いけ Threads フォロワー数", "stock"),
    "note_views": ("いけ note 閲覧数（累計）", "stock"),
    "line_friends": ("いけ 公式LINE 友だち数", "stock"),
    "line_reachable": ("いけ 公式LINE 届く人数", "stock"),
    "check_answers": ("いけ 3分チェック回答数", "flow"),
    "consults": ("いけ 「相談」数", "flow"),
    "seated": ("いけ 面談 着座数", "flow"),
    "contracts": ("いけ 成約数", "flow"),
    "sales": ("いけ 売上", "flow"),
}

# 3分チェックで数えない回答（池田さんの試し）
FORM_NAME_QUESTION = "e94c6cf7-0562-4a89-af76-6a245cdc66de"
FORM_EXCLUDE_NAMES = {"池田和喜"}

JST = dt.timezone(dt.timedelta(hours=9))


def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 80000:  # シリアル値
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    if isinstance(v, str):
        m = re.match(r"^\s*(\d{4})[/-](\d{1,2})[/-](\d{1,2})", v)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
    return None


def num(v):
    if isinstance(v, bool) or v is None:
        return 0
    if isinstance(v, (int, float)):
        return v
    s = re.sub(r"[,，円¥\s]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return 0


def read_dash(path, flow, stock):
    k = json.load(open(os.path.join(path, "kpi_data_A.json")))
    for p in k.get("posts", []):
        d = to_date(p.get("date"))
        if not d:
            continue
        flow["threads_views"][d] += p.get("imp") or 0
        flow["threads_posts"][d] += 1
    for f in k.get("followers", []):
        d = to_date(f.get("date"))
        if d and f.get("count") is not None:
            stock["threads_followers"][d] = f["count"]

    line = json.load(open(os.path.join(path, "line", "line_A.json")))
    for h in line.get("history", []):
        d = to_date(h.get("date"))
        if not d:
            continue
        if h.get("followers") is not None:
            stock["line_friends"][d] = h["followers"]
        if h.get("targetedReaches") is not None:
            stock["line_reachable"][d] = h["targetedReaches"]

    note_path = os.path.join(path, "note", "note_data.json")
    if os.path.exists(note_path):
        n = json.load(open(note_path))
        for day, v in (n.get("history") or {}).items():
            d = to_date(day)
            if d and v.get("pv") is not None:
                stock["note_views"][d] = v["pv"]


def read_sheet(path, flow):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["記録"]
    head = [str(c.value or "").strip() for c in ws[1]]

    def col(name):
        for i, h in enumerate(head):
            if h.startswith(name):
                return i
        raise SystemExit(f"記録タブに列「{name}」が見つからない: {head}")

    c_consult, c_meet, c_result = col("「相談」"), col("面談日"), col("面談の結果")
    c_close, c_sales = col("成約日"), col("売上")
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v not in (None, "") for v in row):
            continue
        d = to_date(row[c_consult])
        if d:
            flow["consults"][d] += 1
        d = to_date(row[c_meet])
        if d and str(row[c_result] or "").strip() == "着座":
            flow["seated"][d] += 1
        d = to_date(row[c_close])
        if d:
            flow["contracts"][d] += 1
            flow["sales"][d] += num(row[c_sales])


def read_form(path, flow):
    data = json.load(open(path))
    if isinstance(data, dict):
        data = data.get("responses") or data.get("items") or []
    for r in data:
        name = ""
        for a in r.get("answers", []):
            if a.get("questionId") == FORM_NAME_QUESTION:
                name = str(a.get("value") or "").strip()
        if name in FORM_EXCLUDE_NAMES:
            continue
        ts = r.get("submittedAt")
        if not ts:
            continue
        d = dt.datetime.fromisoformat(ts.replace("Z", "+00:00")).astimezone(JST).date()
        flow["check_answers"][d] += 1


def build_points(flow, stock, d_from, d_to, has):
    days = [d_from + dt.timedelta(n) for n in range((d_to - d_from).days + 1)]
    out = {}
    for key, (_, kind) in METRICS.items():
        if not has.get(key):
            continue
        pts = []
        if kind == "flow":
            for d in days:
                pts.append({"period": d.isoformat(), "value": flow[key][d]})
        else:
            for d in days:
                if d in stock[key]:
                    pts.append({"period": d.isoformat(), "value": stock[key][d]})
        for p in pts:
            v = p["value"]
            p["value"] = int(v) if float(v).is_integer() else round(v, 4)
        out[key] = pts
    return out


def month_span(d):
    last = calendar.monthrange(d.year, d.month)[1]
    return d.replace(day=1), d.replace(day=last)


def write_points(points, outdir):
    import pyarrow as pa
    import pyarrow.parquet as pq

    os.makedirs(outdir, exist_ok=True)
    summary = {}
    for key, pts in points.items():
        with open(os.path.join(outdir, f"{key}.json"), "w") as f:
            json.dump(pts, f, ensure_ascii=False)
        by_month = collections.defaultdict(list)
        for p in pts:
            by_month[p["period"][:7]].append(p)
        evidence = []
        for ym, rows in sorted(by_month.items()):
            m_from, m_to = month_span(dt.date.fromisoformat(rows[0]["period"]))
            table = pa.table({
                "period": [p["period"] for p in rows],
                "metric": [key] * len(rows),
                "value": [float(p["value"]) for p in rows],
            })
            name = f"{key}_{ym}.parquet"
            path = os.path.join(outdir, name)
            pq.write_table(table, path)
            data = open(path, "rb").read()
            evidence.append({
                "period_from": m_from.isoformat(),
                "period_to": m_to.isoformat(),
                "file": name,
                "size_bytes": len(data),
                "checksum": hashlib.sha256(data).hexdigest(),
                "rows": table.num_rows,
            })
        summary[key] = {"title": METRICS[key][0], "points": len(pts), "evidence": evidence}
    with open(os.path.join(outdir, "summary.json"), "w") as f:
        json.dump(summary, f, ensure_ascii=False, indent=1)
    return summary


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dash", required=True, help="kpi-dashboard の clone")
    ap.add_argument("--sheet", help="「【いけ】属人ルート 面談・成約」の xlsx")
    ap.add_argument("--form-json", help="3分チェックの回答（list_form_responses の responses）")
    ap.add_argument("--from", dest="d_from", required=True)
    ap.add_argument("--to", dest="d_to", required=True)
    ap.add_argument("--points", help="点と証憑を書き出すフォルダ")
    a = ap.parse_args()

    flow = collections.defaultdict(collections.Counter)
    stock = collections.defaultdict(dict)
    read_dash(a.dash, flow, stock)
    has = {k: True for k in ("threads_views", "threads_posts", "threads_followers",
                             "note_views", "line_friends", "line_reachable")}
    if a.sheet:
        read_sheet(a.sheet, flow)
        has.update(consults=True, seated=True, contracts=True, sales=True)
    if a.form_json:
        read_form(a.form_json, flow)
        has["check_answers"] = True

    d_from, d_to = dt.date.fromisoformat(a.d_from), dt.date.fromisoformat(a.d_to)
    points = build_points(flow, stock, d_from, d_to, has)

    # 確かめ用：flowは月の合計、stockは月の最後の値
    print("月ごと:")
    for key, pts in points.items():
        title, kind = METRICS[key]
        months = collections.OrderedDict()
        for p in pts:
            m = p["period"][:7]
            months[m] = months.get(m, 0) + p["value"] if kind == "flow" else p["value"]
        print(f"  {title}: " + ", ".join(f"{m} {v}" for m, v in months.items()))
    if a.points:
        s = write_points(points, a.points)
        print(json.dumps({k: v["points"] for k, v in s.items()}, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
