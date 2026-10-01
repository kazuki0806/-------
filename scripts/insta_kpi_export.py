#!/usr/bin/env python3
"""インスタインフルエンサー経由の数字を、日ごとのKPIの点にする。

読むシート（どれも xlsx に書き出したもの）:
  --dm  担当:YYYY-MM=path   「Instagram DM管理 <担当> <月>」 日付タブ(1〜31)のDM数・返信
  --ag  YYYY-MM=path        「【Kir】AGシート（YYYY年M月）」の「インフルエンサーDM」タブ（提携面談）
  --inf YYYY-MM=path        「【Kir】インフルエンサー営業数値（YYYY年M月）」の「営業」「契約者」タブ

出力（--points DIR）:
  <キー>.json     record_analytics_metric_points にそのまま渡す点（合計は毎日・dims なし。
                  内訳は dims 1つで、合計が0でない日だけ）
  <キー>.parquet  証憑（毎日の合計と、内訳の0も含む全表。日付・内訳・値の集計だけで、顧客名は入れない）
  summary.json    キーごとの点の数・size_bytes・checksum・rows
"""
import argparse
import collections
import datetime as dt
import hashlib
import json
import os
import re
import sys

import openpyxl

# キー → (説明, 内訳のキー)
METRICS = {
    "dm_sent": ("インスタ DM送信数", ["sender"]),
    "dm_replies": ("インスタ DM返信数（見込み）", ["sender"]),
    "partner_seated": ("インスタ 提携面談の着席数", ["via"]),
    "partner_deals": ("インスタ 提携決定数", ["via"]),
    "apo": ("インスタ アポ取り数", ["partner"]),
    "seated": ("インスタ 着座数", ["partner"]),
    "contracts": ("インスタ 契約数", ["partner", "course"]),
    "sales": ("インスタ 売上", ["partner", "closer", "course"]),
}


def to_date(v, year=None):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        s = v.strip()
        m = re.match(r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})", s)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        m = re.match(r"^(\d{1,2})/(\d{1,2})", s)
        if m and year:
            return dt.date(year, int(m[1]), int(m[2]))
    return None


def num(v):
    if isinstance(v, bool):
        return 0
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.replace(",", "").replace("¥", "").strip()
        try:
            return float(s)
        except ValueError:
            return 0
    return 0


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v))


def course_of(label):
    s = clean(label)
    for c in ("ライト", "スタンダード", "プレミアム"):
        if c in s:
            return c
    return "その他" if s else "未記入"


def course_price(label):
    """「CWスタンダード(60万)」→ 660000（税込）。読めなければ 0。"""
    m = re.search(r"(\d+)万", clean(label))
    return int(m[1]) * 10000 * 11 // 10 if m else 0


def parse_kv(items, with_name=False):
    out = []
    for it in items or []:
        k, path = it.split("=", 1)
        if with_name:
            name, month = k.split(":", 1)
            out.append((name, month, path))
        else:
            out.append((k, path))
    return out


# ---------- DM管理シート ----------
def read_dm(files, events):
    for name, month, path in files:
        y, m = map(int, month.split("-"))
        wb = openpyxl.load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            if not ws.title.strip().isdigit():
                continue
            try:
                day = dt.date(y, m, int(ws.title.strip()))
            except ValueError:
                continue
            # 列は 新規DM数=C, 返信 見込み=E, 見込み外=F, アポ取り数=G（1〜2行目が見出し）
            for r in range(3, ws.max_row + 1):
                a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
                if a == "合計" or b == "合計":
                    continue
                dm, rep = num(ws.cell(r, 3).value), num(ws.cell(r, 5).value)
                if dm:
                    events["dm_sent"].append((day, {"sender": name}, dm))
                if rep:
                    events["dm_replies"].append((day, {"sender": name}, rep))


# ---------- AGシート「インフルエンサーDM」（提携面談） ----------
def read_ag(files, events):
    rows = {}  # 相手のアカウント → (月, 行)。新しい月のシートを正にする
    for month, path in sorted(files):
        y = int(month[:4])
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if "インフルエンサーDM" not in wb.sheetnames:
            continue
        ws = wb["インフルエンサーDM"]
        header = None
        for row in ws.iter_rows(values_only=True):
            vals = [clean(v) for v in row]
            if header is None:
                if "面談日" in vals:
                    header = {}
                    for i, h in enumerate(vals):
                        header.setdefault(h, i)  # 同じ見出しは最初の列
                continue
            get = lambda h: row[header[h]] if h in header and header[h] < len(row) else None
            day = to_date(get("面談日"), y)
            acct = clean(get("アカウント名"))
            url = clean(get("URL"))
            if not day or not (acct or url):
                continue
            m = re.search(r"instagram\.com/([^/?]+)", url)
            key = (m[1] if m else acct).lower()
            rows[key] = (month, day, clean(get("経由")) or "未記入", get("着席") is True, get("契約") is True)
    for _, day, via, seated, deal in rows.values():
        if seated:
            events["partner_seated"].append((day, {"via": via}, 1))
        if deal:
            events["partner_deals"].append((day, {"via": via}, 1))


# ---------- インフルエンサー営業数値（提携先経由の販売） ----------
def find_cols(ws):
    """営業タブの列を見出しから探す。1行目に「アポ」「プレ」「契約」のまとまり、4行目に各列の見出し。"""
    h4 = {clean(ws.cell(4, c).value): c for c in range(ws.max_column, 0, -1)}
    groups = {}
    for c in range(1, ws.max_column + 1):
        g = clean(ws.cell(1, c).value)
        if g in ("アポ", "プレ", "契約") and g not in groups:
            groups[g] = c
    cols = {k: h4.get(k) for k in ("顧客名", "リード獲得先", "アポインター", "ステータス", "アポ取り日", "コース")}
    if "アポ" in groups:
        g = groups["アポ"]
        cols["アポ実際"] = next((c for c in range(g, g + 4) if clean(ws.cell(4, c).value) == "実際"), None)
    return cols


def read_inf(files, events):
    leads = {}
    contracts = {}
    for month, path in sorted(files):
        y = int(month[:4])
        wb = openpyxl.load_workbook(path, data_only=True)
        if "営業" in wb.sheetnames:
            ws = wb["営業"]
            c = find_cols(ws)
            for r in range(5, ws.max_row + 1):
                v = lambda k: ws.cell(r, c[k]).value if c.get(k) else None
                name = clean(v("顧客名"))
                apo_taken = to_date(v("アポ取り日"), y)
                if not name or not apo_taken:
                    continue
                key = hashlib.sha256(f"{name}|{apo_taken}".encode()).hexdigest()
                leads[key] = {
                    "partner": clean(v("リード獲得先")) or "未記入",
                    "apo_taken": apo_taken,
                    "seated": to_date(v("アポ実際"), y),
                }
        if "契約者" in wb.sheetnames:
            ws = wb["契約者"]
            h = {clean(ws.cell(1, cc).value): cc for cc in range(ws.max_column, 0, -1)}
            for r in range(2, ws.max_row + 1):
                v = lambda k: ws.cell(r, h[k]).value if k in h else None
                name = clean(v("クライアント名"))
                day = to_date(v("契約日"), y)
                if not name or not day or to_date(v("キャンセル日"), y):
                    continue
                paid = num(v("入金額"))
                amount = paid if v("満額入金済") is True and paid > 0 else course_price(v("コース"))
                key = hashlib.sha256(f"{name}|{day}".encode()).hexdigest()
                contracts[key] = {
                    "day": day,
                    "partner": clean(v("リード獲得先")) or "未記入",
                    "closer": clean(v("担当")) or "未記入",
                    "course": course_of(v("コース")),
                    "amount": amount,
                }
    for x in leads.values():
        events["apo"].append((x["apo_taken"], {"partner": x["partner"]}, 1))
        if x["seated"]:
            events["seated"].append((x["seated"], {"partner": x["partner"]}, 1))
    for x in contracts.values():
        events["contracts"].append((x["day"], {"partner": x["partner"], "course": x["course"]}, 1))
        events["sales"].append(
            (x["day"], {"partner": x["partner"], "closer": x["closer"], "course": x["course"]}, x["amount"])
        )


def build_rows(events, d_from, d_to):
    """証憑の表。合計（dims なし）と内訳（dims 1つ）を、期間の毎日について作る。内訳は出てきた名前全部に0も入れる。"""
    days = [d_from + dt.timedelta(n) for n in range((d_to - d_from).days + 1)]
    out = {}
    for key, (_, dim_keys) in METRICS.items():
        total = collections.Counter()
        by = {k: collections.Counter() for k in dim_keys}
        cats = {k: set() for k in dim_keys}
        for day, dims, value in events[key]:
            for k in dim_keys:
                cats[k].add(dims.get(k, "未記入"))
            if not (d_from <= day <= d_to):
                continue
            total[day] += value
            for k in dim_keys:
                by[k][(day, dims.get(k, "未記入"))] += value
        pts = []
        for day in days:
            pts.append({"period": day.isoformat(), "value": total[day]})
            for k in dim_keys:
                for cat in sorted(cats[k]):
                    pts.append({"period": day.isoformat(), "value": by[k][(day, cat)], "dims": {k: cat}})
        for p in pts:
            p["value"] = int(p["value"]) if float(p["value"]).is_integer() else round(p["value"], 4)
        out[key] = pts
    return out


def sparse(rows):
    """記録する点。合計は毎日入れ、内訳は合計が0でない日だけ入れる（点の数を抑える）。"""
    nonzero = {p["period"] for p in rows if "dims" not in p and p["value"]}
    return [p for p in rows if "dims" not in p or p["period"] in nonzero]


def write_points(rows, outdir, d_from, d_to):
    import pyarrow as pa
    import pyarrow.parquet as pq

    os.makedirs(outdir, exist_ok=True)
    summary = {}
    for key, pts in rows.items():
        rec = sparse(pts)
        with open(os.path.join(outdir, f"{key}.json"), "w") as f:
            json.dump(rec, f, ensure_ascii=False)
        table = pa.table({
            "period": [p["period"] for p in pts],
            "dim_key": [next(iter(p["dims"])) if "dims" in p else "" for p in pts],
            "dim_value": [next(iter(p["dims"].values())) if "dims" in p else "" for p in pts],
            "value": [float(p["value"]) for p in pts],
        })
        path = os.path.join(outdir, f"{key}.parquet")
        pq.write_table(table, path)
        data = open(path, "rb").read()
        summary[key] = {
            "points": len(rec),
            "rows": table.num_rows,
            "size_bytes": len(data),
            "checksum": hashlib.sha256(data).hexdigest(),
            "period_from": d_from.isoformat(),
            "period_to": d_to.isoformat(),
        }
    with open(os.path.join(outdir, "summary.json"), "w") as f:
        json.dump(summary, f, ensure_ascii=False, indent=1)
    return summary


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dm", action="append", help="担当:YYYY-MM=path")
    ap.add_argument("--ag", action="append", help="YYYY-MM=path")
    ap.add_argument("--inf", action="append", help="YYYY-MM=path")
    ap.add_argument("--from", dest="d_from", required=True)
    ap.add_argument("--to", dest="d_to", required=True)
    ap.add_argument("--points", help="点と証憑を書き出すフォルダ")
    a = ap.parse_args()

    events = collections.defaultdict(list)
    read_dm(parse_kv(a.dm, with_name=True), events)
    read_ag(parse_kv(a.ag), events)
    read_inf(parse_kv(a.inf), events)

    d_from, d_to = dt.date.fromisoformat(a.d_from), dt.date.fromisoformat(a.d_to)
    points = build_rows(events, d_from, d_to)

    # 月ごとの合計（確かめ用）
    print("月ごとの合計:")
    for key, (title, _) in METRICS.items():
        months = collections.Counter()
        for p in points[key]:
            if "dims" not in p:
                months[p["period"][:7]] += p["value"]
        print(f"  {title}: " + ", ".join(f"{m} {int(v) if float(v).is_integer() else v}" for m, v in sorted(months.items())))
    if a.points:
        s = write_points(points, a.points, d_from, d_to)
        print(json.dumps({k: v["points"] for k, v in s.items()}, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
