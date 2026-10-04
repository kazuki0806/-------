#!/usr/bin/env python3
"""threads 案件掲載アカウント経由（マーケ別）の数字を、日ごとのKPIの点にする。

読むシート（Google Drive の「【Kir】AGシート（YYYY年M月）」を xlsx に書き出したもの）:
  --ag YYYY-MM=path  複数可。「営業(threads)」タブ（1件1行の台帳。4行目が見出し、5行目からデータ）と
                     「契約者」タブ（1行目が見出し）を読む。
                     同じ人が月をまたいで両方のシートに載っている時は、新しい月のシートの行を正にする
                     （行の同一性は 顧客名＋アポ取り日 のハッシュ。名前そのものは残さない）。

数えるもの（キー → KPI名。内訳 marketer は台帳G列「リード獲得先」、空欄は「未記入」）:
  apo          threads アポ取り数        … 台帳の行（アポ取り日。無ければトスアップ予定日）
  noshow       threads トスアップ飛び数   … ステータス「トスアップ飛び」（トスアップ予定日）
  seated       threads 着座数            … トスアップ「実際」に日付がある行（その日）。見込み外も数える
  unqualified  threads 見込み外数        … ステータスが「見込み外(…)」の行（着座日）。内訳 reason は括弧の中
  closer_apo   threads クローザーアポ数   … 着座し、アポ（クローザー）「予定」に日付がある行（着座日）
  contracts    threads 契約数            … 契約者タブ・種別に threads を含む・キャンセル日なし（契約日）
                                           kpi-dashboard の sales_kpi_export.py（threads_contracts）と同じ

出力（--points DIR）:
  <キー>.json             record_analytics_metric_points にそのまま渡す点
                          （合計は毎日・dims なし。内訳は dims 1つで、合計が0でない日だけ）
  <キー>_<YYYY-MM>.parquet 証憑（月ごと。毎日の合計と内訳の0も含む全表。日付・内訳・値だけで、顧客名は入れない）
  summary.json            キーごとの点の数と evidence（period_from・period_to・file・size_bytes・checksum・rows）
                          証憑の期間は月の1日〜末日で固定（Addness は部分的に重なる証憑を受け付けないため）
"""
import argparse
import calendar
import collections
import datetime as dt
import hashlib
import json
import os
import re
import sys

import openpyxl

METRICS = {
    "apo": ("threads アポ取り数", ["marketer"]),
    "noshow": ("threads トスアップ飛び数", ["marketer"]),
    "seated": ("threads 着座数", ["marketer"]),
    "unqualified": ("threads 見込み外数", ["marketer", "reason"]),
    "closer_apo": ("threads クローザーアポ数", ["marketer"]),
    "contracts": ("threads 契約数", ["marketer"]),
}
BLANK = "未記入"
LEDGER_TAB = "営業(threads)"
CONTRACT_TAB = "契約者"


def to_date(v, year=None):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str):
        s = v.strip()
        m = re.match(r"^(\d{4})[/\-.年](\d{1,2})[/\-.月](\d{1,2})", s)
        if m:
            return dt.date(int(m[1]), int(m[2]), int(m[3]))
        m = re.match(r"^(\d{1,2})[/月](\d{1,2})", s)
        if m and year:
            return dt.date(year, int(m[1]), int(m[2]))
    return None


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v))


def key_of(*parts):
    return hashlib.sha256("|".join(str(p) for p in parts).encode()).hexdigest()


def reason_of(status):
    """「見込み外(支払い能力なし)」→「支払い能力なし」。括弧が無ければ「理由なし」。"""
    m = re.search(r"[（(](.+?)[)）]", status)
    return clean(m[1]) if m else "理由なし"


def parse_kv(items):
    out = []
    for it in items or []:
        k, path = it.split("=", 1)
        out.append((k, path))
    return out


# ---------- 営業(threads) ----------
def ledger_cols(ws):
    """1行目の段（トスアップ・アポ・プレ・契約）と4行目の見出しから列を探す（列の位置が月で変わっても動く）。"""
    head = list(ws.iter_rows(min_row=1, max_row=4, values_only=True))
    if len(head) < 4:
        raise SystemExit(f"{LEDGER_TAB}: 見出しの行（1〜4行目）が足りません")
    row1 = {i + 1: clean(v) for i, v in enumerate(head[0])}
    row4 = {i + 1: clean(v) for i, v in enumerate(head[3])}
    by_name = {}
    for c, h in row4.items():
        by_name.setdefault(h, c)
    cols = {
        "name": by_name.get("顧客名"),
        "marketer": by_name.get("リード獲得先"),
        "status": by_name.get("ステータス"),
        "apo_taken": by_name.get("アポ取り日"),
        "remarks": by_name.get("備考欄"),
    }
    groups = {}
    for c, g in row1.items():
        if g in ("トスアップ", "アポ", "プレ", "契約") and g not in groups:
            groups[g] = c

    def within(g, label):
        start = groups.get(g)
        if not start:
            return None
        return next((c for c in range(start, start + 4) if row4.get(c) == label), None)

    cols["tu_plan"] = within("トスアップ", "予定")
    cols["tu_done"] = within("トスアップ", "実際")
    cols["apo_plan"] = within("アポ", "予定")
    missing = [k for k, v in cols.items() if v is None and k != "remarks"]
    if missing:
        raise SystemExit(f"{LEDGER_TAB}: 列が見つかりません: {missing}")
    return cols


def read_ledgers(files):
    rows = {}  # 行キー → (月, 行の要約)。新しい月のシートを正にする
    for month, path in sorted(files):
        y = int(month[:4])
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if LEDGER_TAB not in wb.sheetnames:
            print(f"{path}: {LEDGER_TAB} タブが無い", file=sys.stderr)
            continue
        ws = wb[LEDGER_TAB]
        c = ledger_cols(ws)
        for r in ws.iter_rows(min_row=5, values_only=True):
            v = lambda k: r[c[k] - 1] if c.get(k) and c[k] - 1 < len(r) else None
            name = clean(v("name"))
            apo_taken = to_date(v("apo_taken"), y)
            tu_plan = to_date(v("tu_plan"), y)
            if not name and not apo_taken:
                continue
            status = clean(v("status"))
            rows[key_of(name, apo_taken or tu_plan)] = {
                "month": month,
                "marketer": clean(v("marketer")) or BLANK,
                "status": status,
                "apo_taken": apo_taken,
                "tu_plan": tu_plan,
                "tu_done": to_date(v("tu_done"), y),
                "apo_plan": to_date(v("apo_plan"), y),
                "remarks": bool(clean(v("remarks"))),
            }
    return rows


# ---------- 契約者 ----------
def read_contracts(files):
    out = {}
    for month, path in sorted(files):
        y = int(month[:4])
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if CONTRACT_TAB not in wb.sheetnames:
            continue
        ws = wb[CONTRACT_TAB]
        hdr = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        h = {}
        for i, x in enumerate(hdr):
            h.setdefault(clean(x), i)
        need = ["クライアント名", "種別", "契約日", "キャンセル日", "リード獲得先"]
        miss = [k for k in need if k not in h]
        if miss:
            print(f"{path}: {CONTRACT_TAB} の列が無い: {miss}", file=sys.stderr)
            continue
        for r in ws.iter_rows(min_row=2, values_only=True):
            v = lambda k: r[h[k]] if h[k] < len(r) else None
            kind = clean(v("種別")).lower()
            day = to_date(v("契約日"), y)
            if "threads" not in kind or not day or clean(v("キャンセル日")):
                continue
            out[key_of(clean(v("クライアント名")), day)] = {
                "day": day,
                "marketer": clean(v("リード獲得先")) or BLANK,
            }
    return out


def make_events(ledger, contracts):
    ev = collections.defaultdict(list)
    for x in ledger.values():
        mk = {"marketer": x["marketer"]}
        apo_day = x["apo_taken"] or x["tu_plan"]
        if apo_day:
            ev["apo"].append((apo_day, mk, 1))
        if x["status"] == "トスアップ飛び" and (x["tu_plan"] or apo_day):
            ev["noshow"].append((x["tu_plan"] or apo_day, mk, 1))
        if x["tu_done"]:
            ev["seated"].append((x["tu_done"], mk, 1))
            if x["apo_plan"]:
                ev["closer_apo"].append((x["tu_done"], mk, 1))
        if x["status"].startswith("見込み外"):
            day = x["tu_done"] or x["tu_plan"] or apo_day
            if day:
                ev["unqualified"].append((day, {**mk, "reason": reason_of(x["status"])}, 1))
    for x in contracts.values():
        ev["contracts"].append((x["day"], {"marketer": x["marketer"]}, 1))
    return ev


def build_rows(events, d_from, d_to):
    """合計（dims なし）と内訳（dims 1つ）を、期間の毎日について作る。内訳は出てきた名前全部に0も入れる。"""
    days = [d_from + dt.timedelta(n) for n in range((d_to - d_from).days + 1)]
    out = {}
    for key, (_, dim_keys) in METRICS.items():
        total = collections.Counter()
        by = {k: collections.Counter() for k in dim_keys}
        cats = {k: set() for k in dim_keys}
        for day, dims, value in events[key]:
            if not (d_from <= day <= d_to):
                continue
            for k in dim_keys:
                cats[k].add(dims.get(k, BLANK))
            total[day] += value
            for k in dim_keys:
                by[k][(day, dims.get(k, BLANK))] += value
        pts = []
        for day in days:
            pts.append({"period": day.isoformat(), "value": int(total[day])})
            for k in dim_keys:
                for cat in sorted(cats[k]):
                    pts.append({"period": day.isoformat(), "value": int(by[k][(day, cat)]), "dims": {k: cat}})
        out[key] = pts
    return out


def sparse(rows):
    """記録する点。合計は毎日入れ、内訳は合計が0でない日だけ入れる。"""
    nonzero = {p["period"] for p in rows if "dims" not in p and p["value"]}
    return [p for p in rows if "dims" not in p or p["period"] in nonzero]


def month_blocks(d_from, d_to):
    cur = d_from.replace(day=1)
    while cur <= d_to:
        last = cur.replace(day=calendar.monthrange(cur.year, cur.month)[1])
        yield cur, last
        cur = last + dt.timedelta(days=1)


def write_points(rows, outdir, d_from, d_to):
    import pyarrow as pa
    import pyarrow.parquet as pq

    os.makedirs(outdir, exist_ok=True)
    summary = {}
    for key, pts in rows.items():
        rec = sparse(pts)
        with open(os.path.join(outdir, f"{key}.json"), "w") as f:
            json.dump(rec, f, ensure_ascii=False)
        evidence = []
        for m_from, m_to in month_blocks(d_from, d_to):
            ym = m_from.strftime("%Y-%m")
            mp = [p for p in pts if p["period"][:7] == ym]
            table = pa.table({
                "period": [p["period"] for p in mp],
                "dim_key": [next(iter(p["dims"])) if "dims" in p else "" for p in mp],
                "dim_value": [next(iter(p["dims"].values())) if "dims" in p else "" for p in mp],
                "value": [float(p["value"]) for p in mp],
            })
            fname = f"{key}_{ym}.parquet"
            path = os.path.join(outdir, fname)
            pq.write_table(table, path)
            data = open(path, "rb").read()
            evidence.append({
                "period_from": m_from.isoformat(),
                "period_to": m_to.isoformat(),
                "file": fname,
                "rows": table.num_rows,
                "size_bytes": len(data),
                "checksum": hashlib.sha256(data).hexdigest(),
            })
        summary[key] = {"title": METRICS[key][0], "points": len(rec), "evidence": evidence}
    with open(os.path.join(outdir, "summary.json"), "w") as f:
        json.dump(summary, f, ensure_ascii=False, indent=1)
    return summary


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ag", action="append", required=True, help="YYYY-MM=path")
    ap.add_argument("--from", dest="d_from", required=True)
    ap.add_argument("--to", dest="d_to", required=True)
    ap.add_argument("--points", help="点と証憑を書き出すフォルダ（リポジトリの外にする）")
    ap.add_argument("--by-marketer", action="store_true", help="確かめ用に、月ごとのマーケ別の件数も表示する")
    a = ap.parse_args()

    files = parse_kv(a.ag)
    ledger = read_ledgers(files)
    contracts = read_contracts(files)
    events = make_events(ledger, contracts)

    d_from, d_to = dt.date.fromisoformat(a.d_from), dt.date.fromisoformat(a.d_to)
    points = build_rows(events, d_from, d_to)

    print("月ごとの合計:")
    for key, (title, _) in METRICS.items():
        months = collections.Counter()
        for p in points[key]:
            if "dims" not in p:
                months[p["period"][:7]] += p["value"]
        print(f"  {title}: " + ", ".join(f"{m} {v}" for m, v in sorted(months.items())))
    seated = [x for x in ledger.values() if x["tu_done"] and d_from <= x["tu_done"] <= d_to]
    print(f"着座した人の備考欄: 記入 {sum(1 for x in seated if x['remarks'])} / {len(seated)}")
    if a.by_marketer:
        for key, (title, dim_keys) in METRICS.items():
            for dk in dim_keys:
                c = collections.Counter()
                for p in points[key]:
                    if "dims" in p and dk in p["dims"] and p["value"]:
                        c[(p["period"][:7], p["dims"][dk])] += p["value"]
                if c:
                    print(f"  {title} / {dk}: " + ", ".join(f"{m} {v}={n}" for (m, v), n in sorted(c.items())))
    if a.points:
        s = write_points(points, a.points, d_from, d_to)
        print(json.dumps({k: v["points"] for k, v in s.items()}, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
